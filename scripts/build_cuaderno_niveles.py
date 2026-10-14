#!/usr/bin/env python3
"""Cuaderno del profesor (xlsx) para una materia con evaluación DESCRIPTIVA por niveles (sin pesos ni notas numéricas).

Uso: python build_cuaderno_niveles.py <plan.json> --curriculo curriculo_oficial.json --out Cuaderno_profesor.xlsx

Lee de plan.json: centro, docente, materia, curso, horario/calendario (vía sesiones.py), tramos, alumnado.n y
`sdas`: [{n, titulo, tramo, unidades:[{nombre, sesiones:[a,b] (dentro del tramo), criterios:[ids]}]}].
Opcional: plan["fecha_corte"] (por defecto, hoy: marca las sesiones ya transcurridas) y plan["niveles"] (lista de niveles, por defecto IN, SU, BI, NT, SB, los de Canarias; cámbiese según la norma de cada comunidad) y plan["plataformas"]["calificaciones"].

Hojas: Léame, Calendario, Unidades, Sesiones, Alumnado, Asistencia, Eval 1/2/3 (criterio x alumno, lista desplegable
de niveles y recuentos con COUNTIF), Calificación (una por alumno y trimestre; la decide el docente, no hay media).
No incluye datos personales: los alumnos se identifican como "Alumno/a n". Los nombres reales solo en una copia local.
"""
import argparse, datetime as dt, json, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sesiones import sesiones, fmt, d  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

HEAD = PatternFill("solid", fgColor="1F3864")
HFONT = Font(bold=True, color="FFFFFF")
GRIS = PatternFill("solid", fgColor="EDEDED")


def cab(ws, fila, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=fila, column=i, value=c)
        x.fill, x.font = HEAD, HFONT
        x.alignment = Alignment(wrap_text=True, vertical="center")


def ancho(ws, anchos):
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[L(i)].width = w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("plan"); ap.add_argument("--curriculo", required=True); ap.add_argument("--out", required=True)
    a = ap.parse_args()
    plan = json.load(open(a.plan, encoding="utf-8"))
    cur = json.load(open(a.curriculo, encoding="utf-8"))
    niveles = plan.get("niveles", ["IN", "SU", "BI", "NT", "SB"])
    n_al = plan["alumnado"].get("n") or plan["alumnado"]["total"]
    hoy = d(plan["fecha_corte"]) if plan.get("fecha_corte") else dt.date.today()
    crit = {r["id"]: (k, r) for k in cur["competencias"] for r in k["criterios"]}
    ses, _ = sesiones(plan)
    tramos = plan["tramos"]
    efem = {x for x in plan["calendario"].get("efemerides_dias_lectivos", {}).get("fechas", []) if ".." not in x}
    eval_dias = {f for e in plan["calendario"].get("sesiones_evaluacion", []) for f in e["fechas"]}

    wb = Workbook()
    ws = wb.active; ws.title = "Léame"
    txt = [
        f"Cuaderno del profesor · {plan['materia']} · {plan['curso']} · {plan['centro']} · {plan['anio']}",
        f"Docente: {plan['docente']}",
        "",
        "Cómo se usa:",
        f"1. Hoja Sesiones: una fila por sesión real ({len(ses)} en el curso). Anotar la actividad realizada y las observaciones.",
        "2. Hoja Asistencia: A = asiste, F = falta, J = falta justificada, R = retraso.",
        f"3. Hojas Eval 1, Eval 2 y Eval 3: para cada criterio previsto en el trimestre, elegir el nivel de cada alumno ({', '.join(niveles)}). Una celda vacía significa «sin evidencia todavía».",
        "4. Hoja Calificación: la calificación trimestral de cada alumno la decide el docente con el conjunto de evidencias. No se calcula ninguna media ni nota numérica.",
        f"5. Las calificaciones oficiales se registran en {plan.get('plataformas', {}).get('calificaciones', '(plataforma del centro)')}.",
        "",
        f"Modelo de evaluación: {plan.get('evaluacion_modelo', 'descriptivo por niveles')}. Niveles de la hoja: {', '.join(niveles)}.",
        "Los alumnos aparecen como «Alumno/a n». NO escribir nombres reales en la copia que se suba al repositorio.",
        "El reparto de unidades y criterios es una PROPUESTA sin validar (ver Propuesta_reparto_PROPUESTA.md).",
    ]
    for i, t in enumerate(txt, 1):
        ws.cell(row=i, column=1, value=t)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 140

    # Calendario
    ws = wb.create_sheet("Calendario")
    cab(ws, 1, ["Fecha", "Día", "Tipo", "Motivo"])
    cal = plan["calendario"]; fila = 2
    filas = [(d(x["fecha"]), "No lectivo", x["motivo"]) for x in cal["no_lectivos"]]
    filas += [(d(e["fechas"][0]) if e["fechas"] else None, "Sesión de evaluación", e["nombre"] + (" · " + ", ".join(e["fechas"]) if e["fechas"] else " · sin fecha")) for e in cal.get("sesiones_evaluacion", [])]
    for p in cal["periodos_sin_clase"]:
        filas.append((d(p["inicio"]), "Periodo sin clase", f"{p['motivo']} (hasta {fmt(d(p['fin']))})"))
    for f, tipo, motivo in sorted([x for x in filas if x[0]], key=lambda x: x[0]):
        ws.cell(row=fila, column=1, value=f).number_format = "DD/MM/YYYY"
        ws.cell(row=fila, column=2, value=fmt(f)[:3]); ws.cell(row=fila, column=3, value=tipo); ws.cell(row=fila, column=4, value=motivo); fila += 1
    ancho(ws, [14, 8, 22, 80])

    # Unidades
    ws = wb.create_sheet("Unidades")
    cab(ws, 1, ["SdA", "Trimestre", "Unidad / tramo", "Sesiones", "Desde", "Hasta", "Criterios previstos"])
    fila = 2
    dentro = {}
    for t in tramos:
        dentro[t["nombre"]] = [s for s in ses if d(t["inicio"]) <= s["fecha"] <= d(t["fin"])]
    for sd in plan["sdas"]:
        dt_ = dentro[sd["tramo"]]
        for u in sd["unidades"]:
            x, y = u["sesiones"]
            vals = [f"SdA {sd['n']} · {sd['titulo']}", sd["tramo"], u["nombre"], y - x + 1, dt_[x - 1]["fecha"], dt_[y - 1]["fecha"], ", ".join(u["criterios"])]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=fila, column=c, value=v)
                if c in (5, 6): cell.number_format = "DD/MM/YYYY"
            fila += 1
    ancho(ws, [42, 16, 40, 10, 12, 12, 50])

    # Sesiones
    ws = wb.create_sheet("Sesiones")
    cols = ["Nº curso", "Trimestre", "Nº en trimestre", "Fecha", "Día", "Franja", "Min", "SdA", "Unidad / tramo", "Criterios previstos", "Nota de calendario", "Estado", "Actividad realizada", "Observaciones"]
    cab(ws, 1, cols)
    fila = 2; n = 0
    for t in tramos:
        sd = next((s for s in plan["sdas"] if s["tramo"] == t["nombre"]), None)
        for i, s in enumerate(dentro[t["nombre"]], 1):
            n += 1
            u = next((u for u in sd["unidades"] if u["sesiones"][0] <= i <= u["sesiones"][1]), None) if sd else None
            f = s["fecha"].isoformat()
            nota = "Sesiones de evaluación (hay clase; cuenta para el trimestre siguiente)" if f in eval_dias else ("Efemérides (día con clase)" if f in efem else "")
            vals = [n, t["nombre"], i, s["fecha"], fmt(s["fecha"])[:3], f"{s['inicio']}-{s['fin']}", s["minutos"], f"SdA {sd['n']}" if sd else "", u["nombre"] if u else "", ", ".join(u["criterios"]) if u else "", nota, "Transcurrida" if s["fecha"] <= hoy else "Prevista", "", ""]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=fila, column=c, value=v)
                if c == 4: cell.number_format = "DD/MM/YYYY"
                if s["fecha"] <= hoy: cell.fill = GRIS
            fila += 1
    ws.freeze_panes = "A2"
    ancho(ws, [9, 16, 10, 12, 6, 13, 6, 8, 38, 28, 40, 13, 40, 40])
    ws.auto_filter.ref = f"A1:N{fila - 1}"
    n_ses = fila - 2

    # Alumnado
    ws = wb.create_sheet("Alumnado")
    cab(ws, 1, ["Código", "Nombre (solo en copia local)", "Medidas / observaciones"])
    for i in range(1, n_al + 1):
        ws.cell(row=i + 1, column=1, value=f"Alumno/a {i}")
    ancho(ws, [14, 36, 60])

    # Asistencia
    ws = wb.create_sheet("Asistencia")
    cab(ws, 1, ["Nº curso", "Fecha", "Franja"] + [f"Alumno/a {i}" for i in range(1, n_al + 1)])
    fila = 2
    for r in range(2, n_ses + 2):
        ws.cell(row=fila, column=1, value=f"=Sesiones!A{r}")
        c = ws.cell(row=fila, column=2, value=f"=Sesiones!D{r}"); c.number_format = "DD/MM/YYYY"
        ws.cell(row=fila, column=3, value=f"=Sesiones!F{r}")
        fila += 1
    dv = DataValidation(type="list", formula1='"A,F,J,R"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"D2:{L(3 + n_al)}{fila - 1}")
    ws.freeze_panes = "D2"
    ancho(ws, [9, 12, 13] + [12] * n_al)
    ini, fin = 2, fila - 1
    ws.cell(row=fin + 2, column=3, value="Faltas (F+J)")
    ws.cell(row=fin + 3, column=3, value="Retrasos (R)")
    for j in range(n_al):
        col = L(4 + j)
        ws.cell(row=fin + 2, column=4 + j, value=f'=COUNTIF({col}{ini}:{col}{fin},"F")+COUNTIF({col}{ini}:{col}{fin},"J")')
        ws.cell(row=fin + 3, column=4 + j, value=f'=COUNTIF({col}{ini}:{col}{fin},"R")')

    # Evaluación por trimestre
    con_sda = [t for t in tramos if any(s["tramo"] == t["nombre"] for s in plan["sdas"])]
    for ti, t in enumerate(con_sda, 1):
        ws = wb.create_sheet(f"Eval {ti}")
        sd = next(s for s in plan["sdas"] if s["tramo"] == t["nombre"])
        previstos = []
        for u in sd["unidades"]:
            for c in u["criterios"]:
                if c not in previstos: previstos.append(c)
        previstos.sort(key=lambda x: tuple(int(p) for p in x.split(".")))
        cab(ws, 1, ["Criterio", "Competencia específica", "Texto del criterio (literal)", "Unidades"] + [f"Alumno/a {i}" for i in range(1, n_al + 1)])
        for r, c in enumerate(previstos, 2):
            k, rr = crit[c]
            ws.cell(row=r, column=1, value=c)
            ws.cell(row=r, column=2, value=f"CE {k['n']}")
            x = ws.cell(row=r, column=3, value=rr["texto"]); x.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=4, value=", ".join(u["nombre"].split(" · ")[0] for u in sd["unidades"] if c in u["criterios"]))
        ultima = len(previstos) + 1
        dv = DataValidation(type="list", formula1='"' + ",".join(niveles) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"E2:{L(4 + n_al)}{ultima}")
        for j, nv in enumerate(niveles):
            r = ultima + 2 + j
            ws.cell(row=r, column=4, value=f"Criterios en {nv}")
            for a_ in range(n_al):
                col = L(5 + a_)
                ws.cell(row=r, column=5 + a_, value=f'=COUNTIF({col}2:{col}{ultima},"{nv}")')
        r = ultima + 2 + len(niveles)
        ws.cell(row=r, column=4, value="Criterios sin evidencia")
        for a_ in range(n_al):
            col = L(5 + a_)
            ws.cell(row=r, column=5 + a_, value=f"=COUNTBLANK({col}2:{col}{ultima})")
        ws.freeze_panes = "E2"
        ancho(ws, [9, 11, 80, 16] + [11] * n_al)

    # Calificación
    ws = wb.create_sheet("Calificación")
    cab(ws, 1, ["Alumno/a"] + [f"{t['nombre']}" for t in con_sda] + ["Observaciones"])
    for i in range(1, n_al + 1):
        ws.cell(row=i + 1, column=1, value=f"Alumno/a {i}")
    dv = DataValidation(type="list", formula1='"' + ",".join(niveles) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"B2:{L(1 + len(con_sda))}{n_al + 1}")
    ws.cell(row=n_al + 3, column=1, value="La calificación la decide el docente. Esta hoja no calcula medias ni notas numéricas.")
    ancho(ws, [16] + [18] * len(con_sda) + [50])

    wb.save(a.out)
    print(f"-> {a.out} ({n_ses} sesiones, {n_al} alumnos, {len(tramos)} trimestres)")


if __name__ == "__main__":
    main()
